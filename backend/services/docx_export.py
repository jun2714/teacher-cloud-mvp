import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Emu

BG_PATH = Path(__file__).resolve().parent.parent / "assets" / "wordbg.png"
TITLE_COLOR = RGBColor(0x1A, 0x5F, 0xB4)
BODY_COLOR = RGBColor(0x2A, 0x31, 0x40)
FONT_NAME = "微软雅黑"


def markdown_to_docx(markdown: str, title: str = "") -> io.BytesIO:
    from services.text_clean import strip_required_marks

    document = Document()
    _setup_page(document)
    _add_page_background(document, BG_PATH)
    _write_markdown(document, strip_required_marks(markdown or ""))
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


def _setup_page(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(28)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)
    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(12)
    style.font.color.rgb = BODY_COLOR
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)


def _add_page_background(document: Document, image_path: Path) -> None:
    if not image_path.exists():
        return
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=section.page_width, height=section.page_height)
    _float_picture_behind_page(run)


def _float_picture_behind_page(run) -> None:
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        return
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    cx = extent.get("cx") if extent is not None else str(int(Emu(Mm(210))))
    cy = extent.get("cy") if extent is not None else str(int(Emu(Mm(297))))
    doc_id = doc_pr.get("id") if doc_pr is not None else "1"
    doc_name = doc_pr.get("name") if doc_pr is not None else "Background"

    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "0")
    anchor.set("behindDoc", "1")
    anchor.set("locked", "1")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    pos_h_off = OxmlElement("wp:posOffset")
    pos_h_off.text = "0"
    pos_h.append(pos_h_off)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    pos_v_off = OxmlElement("wp:posOffset")
    pos_v_off.text = "0"
    pos_v.append(pos_v_off)
    anchor.append(pos_v)

    new_extent = OxmlElement("wp:extent")
    new_extent.set("cx", cx)
    new_extent.set("cy", cy)
    anchor.append(new_extent)

    effect = OxmlElement("wp:effectExtent")
    for key in ("l", "t", "r", "b"):
        effect.set(key, "0")
    anchor.append(effect)
    anchor.append(OxmlElement("wp:wrapNone"))

    new_doc_pr = OxmlElement("wp:docPr")
    new_doc_pr.set("id", doc_id)
    new_doc_pr.set("name", doc_name)
    anchor.append(new_doc_pr)

    cnv = OxmlElement("wp:cNvGraphicFramePr")
    locks = OxmlElement("a:graphicFrameLocks")
    locks.set("noChangeAspect", "1")
    cnv.append(locks)
    anchor.append(cnv)
    if graphic is not None:
        anchor.append(graphic)

    drawing.replace(inline, anchor)


def _write_markdown(document: Document, markdown: str) -> None:
    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _add_table(document, table_lines)
            continue
        if line.startswith("# "):
            _add_heading(document, line[2:].strip(), level=0)
        elif line.startswith("## "):
            _add_heading(document, line[3:].strip(), level=1)
        elif line.startswith("### "):
            _add_heading(document, line[4:].strip(), level=2)
        elif line.startswith("> "):
            _add_quote(document, line[2:].strip())
        elif re.match(r"^[-*] ", line):
            _add_paragraph(document, line[2:].strip(), bullet=True)
        elif line.strip():
            _add_paragraph(document, line.strip())
        i += 1


def _add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    if level == 0:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run(_plain(text))
        _style_run(run, size=18, bold=True, color=TITLE_COLOR)
        return
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(_plain(text))
        _style_run(run, size=14, bold=True, color=TITLE_COLOR)
        return
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(_plain(text))
    _style_run(run, size=12, bold=True, color=BODY_COLOR)


def _add_quote(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Mm(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(_plain(text))
    _style_run(run, size=10.5, color=RGBColor(0x6E, 0x77, 0x86))


def _add_paragraph(document: Document, text: str, bullet: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    if bullet:
        paragraph.style = document.styles["List Bullet"]
    _add_runs_with_bold(paragraph, text)


def _add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-1])
            _style_run(run, size=12, bold=True)
        else:
            run = paragraph.add_run(part)
            _style_run(run, size=12)


def _add_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        if _is_sep_row(line):
            continue
        rows.append(_split_cells(line))
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            value = row[j] if j < len(row) else ""
            run = paragraph.add_run(_plain(value))
            _style_run(run, size=10, bold=(i == 0))
            if i == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8F1FF")
                shading.set(qn("w:val"), "clear")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph()


def _split_cells(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [cell.strip() for cell in text.split("|")]


def _is_sep_row(line: str) -> bool:
    cells = _split_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "") or "") or cell == "" for cell in cells)


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def _style_run(run, size: float = 12, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT_NAME
    run.font.color.rgb = color or BODY_COLOR
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)
