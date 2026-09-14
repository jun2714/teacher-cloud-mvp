from io import BytesIO
from pathlib import Path

MAX_BYTES = 8 * 1024 * 1024
MAX_CHARS = 12000
TEXT_EXT = {".txt", ".md", ".csv", ".json"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
ALLOWED_EXT = TEXT_EXT | IMAGE_EXT | {".docx", ".pdf"}


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(BytesIO(data))
    parts = []
    for page in reader.pages[:40]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)


def extract_upload_excerpt(uploaded) -> tuple[str, str]:
    name = Path(getattr(uploaded, "name", "") or "未命名文件").name
    ext = Path(name).suffix.lower()
    size = int(getattr(uploaded, "size", 0) or 0)
    if ext not in ALLOWED_EXT:
        raise ValueError("暂不支持该文件类型，请上传 txt、md、docx、pdf、csv、json 或图片")
    if size > MAX_BYTES:
        raise ValueError("文件不能超过 8MB")

    if ext in IMAGE_EXT:
        return name, "（图片文件，当前按文本分析，无法直接识别画面内容。请教师补充图片中的要点。）"

    data = uploaded.read()
    uploaded.seek(0)
    if ext in TEXT_EXT:
        text = _decode_bytes(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        text = _pdf_text(data)

    text = (text or "").strip()
    if not text:
        if ext == ".pdf":
            return name, "（未能从 PDF 抽取到可用文字，可能是扫描件。请改传 Word/文本，或补充需要分析的要点。）"
        return name, "（未能从文件中抽取到可用文字，请教师补充需要分析的要点。）"
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n…（内容过长，已截取前段）"
    return name, text
